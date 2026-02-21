#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Módulo para exportar resultados de carreras en múltiples formatos
src/core/race_tracking/results_exporter.py

Soporta exportación a:
- CSV (compatible con Excel)
- PDF (resultados oficiales para imprimir)
- Excel (análisis avanzado con múltiples hojas)
"""

import csv
import logging
from datetime import datetime
from pathlib import Path
from typing import Optional, List, Dict
from dataclasses import dataclass

logger = logging.getLogger(__name__)


@dataclass
class ExportConfig:
    """Configuración para exportación de resultados"""
    event_name: str = "Carrera"
    event_date: Optional[datetime] = None
    organizer: str = ""
    include_splits: bool = True
    include_categories: bool = True
    include_awards: bool = True


class ResultsExporter:
    """
    Exportador de resultados de carreras

    Permite exportar resultados en múltiples formatos con
    diferentes niveles de detalle.

    Example:
        >>> exporter = ResultsExporter(race_manager)
        >>> config = ExportConfig(event_name="Maratón 2026")
        >>> exporter.export_csv("resultados.csv", distance_id="42k", config=config)
        >>> exporter.export_pdf("resultados.pdf", distance_id="42k", config=config)
    """

    def __init__(self, race_manager):
        """
        Inicializar exportador

        Args:
            race_manager: Instancia de RaceManager con los datos
        """
        self.race_manager = race_manager

    def export_csv(self, filepath: str, distance_id: Optional[str] = None,
                   config: Optional[ExportConfig] = None) -> bool:
        """
        Exportar resultados a CSV

        Args:
            filepath: Ruta del archivo CSV a crear
            distance_id: ID de distancia específica (None = todas)
            config: Configuración de exportación

        Returns:
            bool: True si la exportación fue exitosa
        """
        try:
            config = config or ExportConfig()

            # Obtener distancias a exportar
            distances = self._get_distances_to_export(distance_id)
            if not distances:
                logger.warning("⚠️  No hay distancias para exportar")
                return False

            with open(filepath, 'w', newline='', encoding='utf-8-sig') as csvfile:
                # Crear CSV con separador de punto y coma para Excel
                writer = csv.writer(csvfile, delimiter=';')

                # Escribir encabezado del evento
                writer.writerow([f"Evento: {config.event_name}"])
                if config.event_date:
                    writer.writerow([f"Fecha: {config.event_date.strftime('%d/%m/%Y')}"])
                if config.organizer:
                    writer.writerow([f"Organizador: {config.organizer}"])
                writer.writerow([])  # Línea en blanco

                # Exportar cada distancia
                for distance in distances:
                    self._write_distance_to_csv(writer, distance, config)
                    writer.writerow([])  # Separador entre distancias

            logger.info(f"✅ Resultados exportados a CSV: {filepath}")
            return True

        except Exception as e:
            logger.error(f"❌ Error exportando a CSV: {e}")
            import traceback
            traceback.print_exc()
            return False

    def _write_distance_to_csv(self, writer, distance, config: ExportConfig):
        """Escribir resultados de una distancia al CSV"""
        # Encabezado de la distancia
        writer.writerow([f"=== {distance.name} ==="])
        writer.writerow([])

        # Obtener resultados ordenados
        results = self.race_manager.get_results(distance.distance_id)

        if not results:
            writer.writerow(["Sin resultados aún"])
            return

        # Encabezados de columnas base
        headers = ["Pos", "Dorsal", "Nombre", "Género", "Edad", "Categoría",
                   "Tiempo", "Estado"]

        # Agregar columnas de checkpoints si están habilitados
        if config.include_splits and distance.expected_checkpoints > 0:
            for i in range(1, distance.expected_checkpoints + 1):
                headers.append(f"CP{i}")

        writer.writerow(headers)

        # Escribir resultados
        for result in results:
            athlete = result.athlete

            row = [
                result.position or "-",
                athlete.bib_number or "-",
                athlete.name,
                athlete.gender or "-",
                athlete.age or "-",
                athlete.category or "-",
                result.get_formatted_time() if result.finish_time else "-",
                self._get_status_text(result.status)
            ]

            # Agregar splits si están habilitados
            if config.include_splits and distance.expected_checkpoints > 0:
                for i in range(1, distance.expected_checkpoints + 1):
                    if i in result.splits:
                        split = result.splits[i]
                        hours = int(split.total_seconds() // 3600)
                        minutes = int((split.total_seconds() % 3600) // 60)
                        seconds = int(split.total_seconds() % 60)
                        row.append(f"{hours:02d}:{minutes:02d}:{seconds:02d}")
                    else:
                        row.append("-")

            writer.writerow(row)

        # Agregar sección de categorías de premiación si está habilitado
        if config.include_awards:
            writer.writerow([])
            writer.writerow(["=== Clasificación por Categoría ==="])
            writer.writerow([])

            results_by_award = self.race_manager.get_results_by_award_category(distance.distance_id)

            for category_name, category_results in results_by_award.items():
                writer.writerow([f"Categoría: {category_name}"])
                writer.writerow(["Pos", "Dorsal", "Nombre", "Tiempo"])

                for idx, result in enumerate(category_results[:10], 1):  # Top 10
                    writer.writerow([
                        idx,
                        result.athlete.bib_number or "-",
                        result.athlete.name,
                        result.get_formatted_time() if result.finish_time else "-"
                    ])

                writer.writerow([])

    def export_pdf(self, filepath: str, distance_id: Optional[str] = None,
                   config: Optional[ExportConfig] = None) -> bool:
        """
        Exportar resultados a PDF

        Args:
            filepath: Ruta del archivo PDF a crear
            distance_id: ID de distancia específica (None = todas)
            config: Configuración de exportación

        Returns:
            bool: True si la exportación fue exitosa
        """
        # Intentar importar reportlab
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, letter
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
            from reportlab.lib.enums import TA_CENTER, TA_LEFT
        except ImportError:
            logger.error("❌ reportlab no está instalado. Instala con: pip install reportlab")
            raise ImportError("Librería reportlab no instalada")

        try:

            config = config or ExportConfig()

            # Crear documento PDF
            doc = SimpleDocTemplate(filepath, pagesize=A4,
                                   leftMargin=0.5*inch, rightMargin=0.5*inch,
                                   topMargin=0.5*inch, bottomMargin=0.5*inch)

            # Contenedor de elementos
            elements = []
            styles = getSampleStyleSheet()

            # Estilo personalizado para título
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                textColor=colors.HexColor('#1e40af'),
                spaceAfter=30,
                alignment=TA_CENTER
            )

            # Título del evento
            title = Paragraph(config.event_name, title_style)
            elements.append(title)

            # Información del evento
            if config.event_date or config.organizer:
                info_style = ParagraphStyle('Info', parent=styles['Normal'],
                                           fontSize=12, alignment=TA_CENTER)
                if config.event_date:
                    elements.append(Paragraph(
                        f"Fecha: {config.event_date.strftime('%d/%m/%Y')}",
                        info_style
                    ))
                if config.organizer:
                    elements.append(Paragraph(f"Organizador: {config.organizer}", info_style))
                elements.append(Spacer(1, 0.3*inch))

            # Obtener distancias
            distances = self._get_distances_to_export(distance_id)

            for idx, distance in enumerate(distances):
                if idx > 0:
                    elements.append(PageBreak())

                self._add_distance_to_pdf(elements, distance, config, styles)

            # Generar PDF
            doc.build(elements)

            logger.info(f"✅ Resultados exportados a PDF: {filepath}")
            return True

        except Exception as e:
            logger.error(f"❌ Error exportando a PDF: {e}")
            import traceback
            traceback.print_exc()
            return False

    def _add_distance_to_pdf(self, elements, distance, config, styles):
        """Agregar resultados de una distancia al PDF"""
        from reportlab.lib import colors
        from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.enums import TA_CENTER

        # Título de la distancia
        dist_style = styles['Heading2']
        dist_style.alignment = TA_CENTER
        elements.append(Paragraph(distance.name, dist_style))
        elements.append(Spacer(1, 0.2*inch))

        # Obtener resultados
        results = self.race_manager.get_results(distance.distance_id)

        if not results:
            elements.append(Paragraph("Sin resultados aún", styles['Normal']))
            return

        # Crear tabla de resultados
        # Encabezados
        headers = ["Pos", "Dorsal", "Nombre", "Categoría", "Tiempo"]

        data = [headers]

        # Datos
        for result in results[:50]:  # Limitar a top 50 para que quepa en PDF
            athlete = result.athlete
            data.append([
                str(result.position or "-"),
                str(athlete.bib_number or "-"),
                athlete.name,
                athlete.category or "-",
                result.get_formatted_time() if result.finish_time else "-"
            ])

        # Crear y estilizar tabla
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            # Encabezado
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3b82f6')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),

            # Cuerpo
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f4f6')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ]))

        elements.append(table)

    def export_excel(self, filepath: str, distance_id: Optional[str] = None,
                    config: Optional[ExportConfig] = None) -> bool:
        """
        Exportar resultados a Excel con múltiples hojas

        Args:
            filepath: Ruta del archivo Excel a crear
            distance_id: ID de distancia específica (None = todas)
            config: Configuración de exportación

        Returns:
            bool: True si la exportación fue exitosa
        """
        # Intentar importar openpyxl
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            logger.error("❌ openpyxl no está instalado. Instala con: pip install openpyxl")
            raise ImportError("Librería openpyxl no instalada")

        try:

            config = config or ExportConfig()

            # Crear workbook
            wb = Workbook()
            wb.remove(wb.active)  # Remover hoja por defecto

            # Obtener distancias
            distances = self._get_distances_to_export(distance_id)

            for distance in distances:
                self._add_distance_to_excel(wb, distance, config)

            # Guardar archivo
            wb.save(filepath)

            logger.info(f"✅ Resultados exportados a Excel: {filepath}")
            return True

        except Exception as e:
            logger.error(f"❌ Error exportando a Excel: {e}")
            import traceback
            traceback.print_exc()
            return False

    def _add_distance_to_excel(self, wb, distance, config):
        """Agregar hoja de distancia al Excel"""
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        # Crear hoja (nombre limitado a 31 caracteres)
        sheet_name = distance.name[:31]
        ws = wb.create_sheet(title=sheet_name)

        # Estilos
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Título
        ws['A1'] = distance.name
        ws['A1'].font = Font(size=16, bold=True, color="1E40AF")
        ws.merge_cells('A1:G1')
        ws['A1'].alignment = Alignment(horizontal="center")

        # Encabezados (fila 3)
        headers = ["Pos", "Dorsal", "Nombre", "Género", "Edad", "Categoría", "Tiempo"]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        # Obtener resultados
        results = self.race_manager.get_results(distance.distance_id)

        # Datos
        for row_idx, result in enumerate(results, 4):
            athlete = result.athlete

            ws.cell(row=row_idx, column=1, value=result.position or "-")
            ws.cell(row=row_idx, column=2, value=athlete.bib_number or "-")
            ws.cell(row=row_idx, column=3, value=athlete.name)
            ws.cell(row=row_idx, column=4, value=athlete.gender or "-")
            ws.cell(row=row_idx, column=5, value=athlete.age or "-")
            ws.cell(row=row_idx, column=6, value=athlete.category or "-")
            ws.cell(row=row_idx, column=7, value=result.get_formatted_time() if result.finish_time else "-")

            # Aplicar bordes
            for col in range(1, 8):
                ws.cell(row=row_idx, column=col).border = border

        # Ajustar anchos de columna
        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 10
        ws.column_dimensions['E'].width = 8
        ws.column_dimensions['F'].width = 20
        ws.column_dimensions['G'].width = 12

    def _get_distances_to_export(self, distance_id: Optional[str]) -> List:
        """Obtener lista de distancias a exportar"""
        if distance_id:
            distance = self.race_manager.get_distance(distance_id)
            return [distance] if distance else []
        else:
            return self.race_manager.get_all_distances()

    def _get_status_text(self, status) -> str:
        """Convertir estado a texto legible"""
        from .models import AthleteStatus

        status_map = {
            AthleteStatus.NOT_STARTED: "No iniciado",
            AthleteStatus.RUNNING: "En carrera",
            AthleteStatus.FINISHED: "Finalizado",
            AthleteStatus.DNF: "Abandono",
            AthleteStatus.DNS: "No largó",
            AthleteStatus.DQ: "Descalificado"
        }

        return status_map.get(status, str(status))
